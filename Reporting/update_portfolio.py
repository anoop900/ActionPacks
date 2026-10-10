import sys
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

AUTO_FILE = "Retirement_Portfolio_Auto.xlsx"


def D(value, default="0"):
    try:
        if value is None or value == "":
            return Decimal(default)
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal(default)


def as_rate(value):
    rate = D(value)
    if rate > 1 or rate < -1:
        return rate / Decimal("100")
    return rate


def normalize_header(value):
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def header_map(ws):
    return {
        normalize_header(ws.cell(1, col).value): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(1, col).value is not None
    }


def get_decimal(ws, row, headers, *names, fallback=None, default="0"):
    for name in names:
        if name in headers:
            return D(ws.cell(row, headers[name]).value, default)
    if fallback is not None:
        return D(ws.cell(row, fallback).value, default)
    return D(default)


def get_int(ws, row, headers, *names, fallback=None, default="0"):
    return int(get_decimal(ws, row, headers, *names, fallback=fallback, default=default))


def read_cpp_oas(ws):
    headers = header_map(ws)
    cpp_map = {}
    for r in range(2, ws.max_row + 1):
        holder = ws.cell(r, headers.get("holder", 1)).value
        if not holder:
            continue
        cpp_map[holder] = {
            "cpp": get_decimal(ws, r, headers, "cpp", fallback=2),
            "start": get_int(ws, r, headers, "start_age", "start", fallback=3, default="65"),
            "oas": get_decimal(ws, r, headers, "oas", fallback=4),
        }
    return cpp_map


def read_one_time_expenses(ws):
    headers = header_map(ws)
    expense_map = {}
    for r in range(2, ws.max_row + 1):
        holder = ws.cell(r, headers.get("holder", 1)).value
        if not holder:
            continue
        age = get_int(ws, r, headers, "age", fallback=2)
        amt = get_decimal(ws, r, headers, "amount", "expense", fallback=3)
        expense_map.setdefault(holder, []).append((age, amt))
    return expense_map


def read_recurring_expenses(ws):
    if ws is None:
        return {}
    headers = header_map(ws)
    recurring = {}
    for r in range(2, ws.max_row + 1):
        holder = ws.cell(r, headers.get("holder", 1)).value
        if not holder:
            continue
        start_age = get_int(ws, r, headers, "start_age", "start", fallback=2)
        end_age = get_int(ws, r, headers, "end_age", "end", fallback=3)
        amt = get_decimal(ws, r, headers, "annual_amount", "amount", fallback=4)
        recurring.setdefault(holder, []).append((start_age, end_age, amt))
    return recurring


def read_settings(ws):
    headers = header_map(ws)
    cfg = {}
    for r in range(2, ws.max_row + 1):
        holder = ws.cell(r, headers.get("holder", 1)).value
        if not holder:
            continue
        cfg[holder] = {
            "current_age": get_int(ws, r, headers, "current_age", fallback=2),
            "retire_age": get_int(ws, r, headers, "retire_age", fallback=3),
            "life": get_int(ws, r, headers, "life_expectancy", "life", fallback=4),
            "contrib": get_decimal(ws, r, headers, "annual_contribution", "contrib", fallback=5),
            "return": as_rate(get_decimal(ws, r, headers, "annual_return", "return", fallback=6)),
            "tax": as_rate(get_decimal(ws, r, headers, "tax_rate", "tax", fallback=7)),
            "dividend": as_rate(get_decimal(ws, r, headers, "dividend_yield", "dividend", fallback=8, default="0")),
        }
    return cfg


def read_networth_snapshot(ws):
    last = ws.max_row
    prev = last - 1
    year_ago = last - 12
    ag = D(ws.cell(last, 2).value)
    kg = D(ws.cell(last, 3).value)
    ag_prev = D(ws.cell(prev, 2).value) if prev >= 2 else Decimal("0")
    kg_prev = D(ws.cell(prev, 3).value) if prev >= 2 else Decimal("0")
    ag_year = D(ws.cell(year_ago, 2).value) if year_ago >= 2 else Decimal("0")
    kg_year = D(ws.cell(year_ago, 3).value) if year_ago >= 2 else Decimal("0")
    return {
        "AG": {"current": ag, "monthly_growth": ag - ag_prev, "yearly_growth": ag - ag_year},
        "KG": {"current": kg, "monthly_growth": kg - kg_prev, "yearly_growth": kg - kg_year},
        "Combined": {
            "current": ag + kg,
            "monthly_growth": (ag + kg) - (ag_prev + kg_prev),
            "yearly_growth": (ag + kg) - (ag_year + kg_year),
        },
    }


def run_retirement_engine():
    wb = load_workbook(AUTO_FILE)

    snap = wb["NetWorth_History"]
    settings = wb["Retirement_Settings"]
    cpp_ws = wb["CPP_OAS_Settings"]
    expenses_ws = wb["OneTime_Expenses"]
    recurring_ws = wb["Recurring_Expenses"] if "Recurring_Expenses" in wb.sheetnames else None

    for name in ["Retirement_Projection", "Retirement_Summary"]:
        if name in wb.sheetnames:
            del wb[name]

    proj = wb.create_sheet("Retirement_Projection")
    proj.append(
        [
            "Holder",
            "Age",
            "Start_Balance",
            "Investment_Growth",
            "Annual_Contribution",
            "CPP_OAS",
            "Dividend_Income",
            "Recurring_Expenses",
            "OneTime_Expenses",
            "Taxes",
            "End_Balance",
        ]
    )

    summary = wb.create_sheet("Retirement_Summary")
    summary.append(
        [
            "Holder",
            "Current_NetWorth",
            "Monthly_Growth",
            "Yearly_Growth",
            "Dividend_Monthly",
            "Dividend_Yearly",
            "Balance_at_Retirement",
            "Balance_at_Life_Expectancy",
            "Projection_5Y",
            "Projection_10Y",
            "Projection_15Y",
        ]
    )

    networth = read_networth_snapshot(snap)
    cpp_map = read_cpp_oas(cpp_ws)
    expense_map = read_one_time_expenses(expenses_ws)
    recurring_map = read_recurring_expenses(recurring_ws)
    cfg = read_settings(settings)

    results = {}

    for holder in ["AG", "KG"]:
        start_balance = networth[holder]["current"]
        c = cfg[holder]
        bal = start_balance
        bal_at_ret = None
        targets = {c["current_age"] + 5, c["current_age"] + 10, c["current_age"] + 15}
        projections = {}

        for age in range(c["current_age"], c["life"] + 1):
            start = bal
            growth = start * c["return"]
            contrib = c["contrib"] if age < c["retire_age"] else Decimal("0")

            cpp = Decimal("0")
            if holder in cpp_map and age >= cpp_map[holder]["start"]:
                cpp = cpp_map[holder]["cpp"] + cpp_map[holder]["oas"]

            dividend = start * c["dividend"]

            one_time = sum(amt for a, amt in expense_map.get(holder, []) if a == age)
            recurring = sum(
                amt for start_age, end_age, amt in recurring_map.get(holder, []) if start_age <= age <= end_age
            )

            taxable = cpp + dividend
            taxes = taxable * c["tax"] if taxable > 0 else Decimal("0")

            end = start + growth + contrib + cpp - recurring - one_time - taxes

            proj.append(
                [
                    holder,
                    age,
                    float(start),
                    float(growth),
                    float(contrib),
                    float(cpp),
                    float(dividend),
                    float(recurring),
                    float(one_time),
                    float(taxes),
                    float(end),
                ]
            )

            if age == c["retire_age"]:
                bal_at_ret = end

            if age in targets:
                projections[age] = end

            bal = end

        results[holder] = {
            "ret": bal_at_ret if bal_at_ret is not None else start_balance,
            "end": bal,
            "proj_5": projections.get(c["current_age"] + 5, bal),
            "proj_10": projections.get(c["current_age"] + 10, bal),
            "proj_15": projections.get(c["current_age"] + 15, bal),
        }

    for holder in ["AG", "KG", "Combined"]:
        current = networth[holder]["current"]
        monthly_growth = networth[holder]["monthly_growth"]
        yearly_growth = networth[holder]["yearly_growth"]
        if holder == "Combined":
            dividend_yearly = (
                networth["AG"]["current"] * cfg["AG"]["dividend"]
                + networth["KG"]["current"] * cfg["KG"]["dividend"]
            )
            ret = results["AG"]["ret"] + results["KG"]["ret"]
            end = results["AG"]["end"] + results["KG"]["end"]
            proj_5 = results["AG"]["proj_5"] + results["KG"]["proj_5"]
            proj_10 = results["AG"]["proj_10"] + results["KG"]["proj_10"]
            proj_15 = results["AG"]["proj_15"] + results["KG"]["proj_15"]
        else:
            dividend_yearly = current * cfg[holder]["dividend"]
            ret = results[holder]["ret"]
            end = results[holder]["end"]
            proj_5 = results[holder]["proj_5"]
            proj_10 = results[holder]["proj_10"]
            proj_15 = results[holder]["proj_15"]

        summary.append(
            [
                holder,
                float(current),
                float(monthly_growth),
                float(yearly_growth),
                float(dividend_yearly / Decimal("12")),
                float(dividend_yearly),
                float(ret),
                float(end),
                float(proj_5),
                float(proj_10),
                float(proj_15),
            ]
        )

    wb.save(AUTO_FILE)
    print("✔ Retirement model rebuilt with growth, dividends, and projections")


def main():
    if "--retirement" in sys.argv:
        run_retirement_engine()
    else:
        print("Usage: py update_portfolio.py --retirement")


if __name__ == "__main__":
    main()
